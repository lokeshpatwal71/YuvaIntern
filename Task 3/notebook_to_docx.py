"""
notebook_to_docx.py - converts the Jupyter notebook into a Word (.docx) document.

    python src/notebook_to_docx.py
    python src/notebook_to_docx.py --notebook real_world_dataset_analysis.ipynb --out real_world_dataset_analysis.docx

What goes into the Word file
  * every Markdown cell (headings, text, tables, lists)         -> formatted text
  * every code cell                                             -> shaded code block
  * every saved output of a code cell (text, tables, charts)    -> output block / picture

IMPORTANT: outputs are only included if they were SAVED in the notebook.
Run the notebook (Kernel > Restart & Run All), press Ctrl+S, and then run this script again.
A notebook that has not been run produces a Word file with explanations and code only.
"""
from __future__ import annotations

import argparse
import base64
import io
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

import build_report as br        # re-use the styling helpers and the student details

PROJECT_ROOT = Path(__file__).resolve().parent.parent
NAVY, BLUE, GREY = br.NAVY, br.BLUE, br.GREY
MAX_OUTPUT_LINES = 60


# ----------------------------- inline Markdown -----------------------------
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text: str, size=None, color=None):
    """Write text with **bold**, `code`, *italic* and [link](url) into a paragraph."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = paragraph.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = paragraph.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt((size or 11) - 1.5)
            r.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
            continue
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = paragraph.add_run(part[1:-1]); r.italic = True
        elif part.startswith("[") and "](" in part:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part).groups()
            r = paragraph.add_run(label)
        else:
            r = paragraph.add_run(part)
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color


def strip_inline(text: str) -> str:
    return re.sub(r"\*\*|`", "", text)


def md_table(doc, rows):
    header, body = rows[0], rows[2:]
    n = len(header)
    weights = []
    for i in range(n):
        cells = [header[i]] + [r[i] if i < len(r) else "" for r in body]
        lens = [len(strip_inline(c)) for c in cells]
        weights.append(min(max(0.5 * max(lens) + 0.5 * sum(lens) / len(lens), 8), 40))
    total = sum(weights)
    widths = [16.4 * w / total for w in weights]
    t = doc.add_table(rows=1, cols=n); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; br.shade(c, "1F3864")
        r = c.paragraphs[0].add_run(strip_inline(h)); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(255, 255, 255)
    for ri, row in enumerate(body):
        cells = t.add_row().cells
        for i in range(n):
            cells[i].text = ""
            add_inline(cells[i].paragraphs[0], row[i] if i < len(row) else "", size=9)
            if ri % 2: br.shade(cells[i], "F2F6FB")
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
        for c in row.cells:
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def quote(doc, text):
    t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"; c = t.rows[0].cells[0]; br.shade(c, "EAF2FB"); c.text = ""
    add_inline(c.paragraphs[0], text, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render_markdown(doc, text: str, skip_h1=True):
    lines = text.split("\n"); i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip() or ln.strip() == "---":
            i += 1; continue
        m = re.match(r"^(#{1,4})\s+(.*)", ln)
        if m:
            level = len(m.group(1)); title = strip_inline(m.group(2)).strip(); i += 1
            if level == 1:
                if not skip_h1:
                    doc.add_heading(title, level=1)
            else:
                doc.add_heading(title, level=min(level - 1, 3))
            continue
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")]); i += 1
            md_table(doc, rows); continue
        if ln.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].startswith(">"):
                buf.append(lines[i].lstrip("> ").strip()); i += 1
            quote(doc, " ".join(buf)); continue
        if re.match(r"^\d+\.\s", ln):
            n = 0
            while i < len(lines) and re.match(r"^\d+\.\s", lines[i]):
                n += 1; p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.9); p.paragraph_format.first_line_indent = Cm(-0.6)
                p.paragraph_format.space_after = Pt(2)
                p.add_run(f"{n}.  "); add_inline(p, re.sub(r"^\d+\.\s", "", lines[i])); i += 1
            continue
        if re.match(r"^[*-]\s", ln):
            while i < len(lines) and re.match(r"^[*-]\s", lines[i]):
                p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(2)
                add_inline(p, re.sub(r"^[*-]\s", "", lines[i])); i += 1
            continue
        buf = []
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,4}\s|\||>|\d+\.\s|[*-]\s|---)", lines[i]):
            buf.append(lines[i].strip()); i += 1
        p = doc.add_paragraph(); add_inline(p, " ".join(buf))


# ----------------------------- code and outputs -----------------------------
def block(doc, text: str, label: str, fill: str, label_color):
    lab = doc.add_paragraph(); lab.paragraph_format.space_after = Pt(1); lab.paragraph_format.keep_with_next = True
    r = lab.add_run(label); r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = label_color
    t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"; c = t.rows[0].cells[0]; br.shade(c, fill); c.text = ""
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text.rstrip("\n")); run.font.name = "Consolas"; run.font.size = Pt(8)
    run._element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "Consolas")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def join(src) -> str:
    return "".join(src) if isinstance(src, list) else str(src)


def render_outputs(doc, outputs) -> int:
    shown = 0
    for out in outputs:
        kind = out.get("output_type")
        if kind == "stream":
            text = join(out["text"])
        elif kind in ("display_data", "execute_result"):
            data = out.get("data", {})
            if "image/png" in data:
                doc.add_picture(io.BytesIO(base64.b64decode(join(data["image/png"]))), width=Cm(15))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER; shown += 1; continue
            if "text/markdown" in data:
                render_markdown(doc, join(data["text/markdown"]), skip_h1=False); shown += 1; continue
            text = join(data.get("text/plain", ""))
        elif kind == "error":
            text = "ERROR: " + out.get("ename", "") + ": " + out.get("evalue", "")
        else:
            continue
        lines = text.split("\n")
        if len(lines) > MAX_OUTPUT_LINES:
            text = "\n".join(lines[:MAX_OUTPUT_LINES]) + f"\n... ({len(lines) - MAX_OUTPUT_LINES} more lines not shown; see the notebook)"
        if text.strip():
            block(doc, text, "Output", "F1F8F1", RGBColor(0x2E, 0x7D, 0x32)); shown += 1
    return shown


# ----------------------------- document -----------------------------
def convert(nb_path: Path, out_path: Path) -> dict:
    nb = json.loads(Path(nb_path).read_text(encoding="utf-8"))
    nb_cells = nb["cells"]
    n_out = sum(len(c.get("outputs", [])) for c in nb_cells if c["cell_type"] == "code")
    doc = Document(); br.style_document(doc)
    sec = doc.sections[0]; sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]; hp.text = ""
    r = hp.add_run("Real-World Dataset Analysis Project"); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = NAVY
    r = hp.add_run("\t\tJupyter Notebook (Word version)"); r.font.size = Pt(9); r.font.color.rgb = GREY
    fp = sec.footer.paragraphs[0]; fp.text = ""
    r = fp.add_run(f"{br.STUDENT}  |  {br.COLLEGE}\t\tPage "); r.font.size = Pt(9); r.font.color.rgb = GREY
    pr = br.add_field(fp, "PAGE"); pr.font.size = Pt(9); pr.font.color.rgb = GREY

    # cover page
    for _ in range(3): doc.add_paragraph()
    br.para(doc, "INTERNSHIP PROJECT", bold=True, size=13, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
    br.para(doc, "Real-World Dataset Analysis", bold=True, size=30, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    br.para(doc, "Jupyter Notebook: Word version", size=15, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)
    br.para(doc, "Python | Pandas | NumPy | SciPy | Matplotlib | Seaborn", size=11, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(3): doc.add_paragraph()
    t = doc.add_table(rows=0, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for a, b in (("Submitted by:", br.STUDENT), ("Course / Year:", br.COURSE), ("College / University:", br.COLLEGE),
                 ("Internship Provider:", br.INTERNSHIP), ("Internship Period:", br.PERIOD), ("Date of Submission:", br.SUBMISSION_DATE)):
        cells = t.add_row().cells
        r = cells[0].paragraphs[0].add_run(a); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = NAVY
        r = cells[1].paragraphs[0].add_run(b); r.bold = True; r.font.size = Pt(12)
        cells[0].width, cells[1].width = Cm(5.5), Cm(9.5)
    doc.add_page_break()

    # status + contents
    if n_out == 0:
        br.note_box(doc, "NOTEBOOK NOT RUN YET - OUTPUTS PENDING.",
                    "This Word file contains the explanations and the code of the notebook. The tables, numbers and charts appear only after you run the notebook "
                    "on the real dataset (Kernel > Restart & Run All), save it (Ctrl+S) and run:  python src/notebook_to_docx.py")
    else:
        br.note_box(doc, "Outputs included.", f"{n_out} saved outputs (text, tables and charts) from the executed notebook are included below.", fill="EDF7EE")
    doc.add_heading("Contents", level=1)
    for c in nb_cells:
        if c["cell_type"] == "markdown":
            for m in re.finditer(r"^##\s+(.*)$", join(c["source"]), re.M):
                br.para(doc, strip_inline(m.group(1)), space_after=3)
    doc.add_page_break()

    first_md = True
    for c in nb_cells:
        src = join(c["source"])
        if c["cell_type"] == "markdown":
            render_markdown(doc, src, skip_h1=first_md); first_md = False
        else:
            block(doc, src, "Code", "F6F8FA", BLUE)
            render_outputs(doc, c.get("outputs", []))
    out_path = Path(out_path); out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return {"code_cells": sum(c["cell_type"] == "code" for c in nb_cells), "markdown_cells": sum(c["cell_type"] == "markdown" for c in nb_cells), "saved_outputs": n_out}


def main(argv=None):
    ap = argparse.ArgumentParser(description="Convert the Jupyter notebook to a Word document")
    ap.add_argument("--notebook", default=str(PROJECT_ROOT / "real_world_dataset_analysis.ipynb"))
    ap.add_argument("--out", default=str(PROJECT_ROOT / "real_world_dataset_analysis.docx"))
    args = ap.parse_args(argv)
    info = convert(Path(args.notebook), Path(args.out))
    print(f"Created {args.out}: {info['markdown_cells']} text cells, {info['code_cells']} code cells, {info['saved_outputs']} saved outputs.")
    if info["saved_outputs"] == 0:
        print("NOTE: the notebook has no saved outputs yet. Run it, save it (Ctrl+S) and run this script again to include results.")


if __name__ == "__main__":
    main()
